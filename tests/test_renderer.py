"""渲染层测试 — docpipeline/renderer.py + agents/renderer_agent.py

覆盖：
- Markdown 解析（标题层级 / 列表 / 围栏代码 / 引用 / 分隔线）
- docx 渲染：Word 标题样式落地（决定能否自动生成目录）
- pdf 渲染：中文保留、代码缩进保真
- 格式分发与不支持格式的降级
- Agent 契约：格式归一化、文件名解析、失败不中断流水线
- pipeline YAML 可被 scheduler 解析（防止声明与 Agent 不一致）

后端缺失时用 monkeypatch 强制 HAS_DOCX/HAS_PDF=False，
断言降级行为（返回 error 而非抛异常）。
"""
import sys
from pathlib import Path

import pytest

from docpipeline import renderer

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SAMPLE_MD = """# 季度技术报告

## 概述

本季度围绕异步编程能力做了系统性重构，核心目标是把检索、写作与渲染
三个环节串成可交付的流水线，并保证产出格式可用于打印与送审。

### 关键进展

- 完成 DAG 编排内核，支持断点续传与熔断保护
- 正文提取保留代码块换行与缩进
- 渲染层打通 docx 与 pdf 双路线

```python
async def render(doc: str) -> bytes:
    result = await pipeline.run(doc)
    return result.to_pdf()
```

> 说明：渲染是增强环节，失败不影响 Markdown 主产物。

---

## 后续计划

下季度补齐知识库摄入侧。
"""

needs_docx = pytest.mark.skipif(not renderer.HAS_DOCX, reason="python-docx 未安装")
needs_pdf = pytest.mark.skipif(not renderer.HAS_PDF, reason="reportlab 未安装")

# pymupdf 仅用于校验 pdf 文本层，是测试期依赖而非运行时依赖；
# 未安装时跳过这几个用例，而不是让整条 CI 变红。
try:
    import pymupdf  # noqa: F401
    HAS_PYMUPDF = True
except ImportError:      # pragma: no cover - 环境相关
    HAS_PYMUPDF = False
needs_pdftext = pytest.mark.skipif(not HAS_PYMUPDF, reason="pymupdf 未安装")


# ───────────────────────────── Markdown 解析 ─────────────────────────────

class TestParseMarkdown:
    def test_heading_levels(self):
        kinds = [k for k, _ in renderer.parse_markdown("# a\n## b\n### c")]
        assert kinds == ["h1", "h2", "h3"]

    def test_list_items(self):
        blocks = renderer.parse_markdown("- 第一项\n* 第二项\n1. 第三项")
        assert [k for k, _ in blocks] == ["li", "li", "li"]
        assert blocks[2][1] == "第三项"

    def test_fenced_code_preserved(self):
        blocks = renderer.parse_markdown("```python\nx = 1\n    y = 2\n```")
        assert blocks[0][0] == "code"
        assert "    y = 2" in blocks[0][1]

    def test_tilde_fence(self):
        blocks = renderer.parse_markdown("~~~\nplain\n~~~")
        assert blocks[0][0] == "code"

    def test_inline_code_stripped(self):
        assert renderer.parse_markdown("调用 `foo()` 完成")[0][1] == "调用 foo() 完成"

    def test_bold_stripped(self):
        assert renderer.parse_markdown("**重点**内容")[0][1] == "重点内容"

    def test_link_keeps_url(self):
        # 引用来源是可信度关键信息，URL 不能丢
        out = renderer.parse_markdown("见 [文档](https://x.com/a)")[0][1]
        assert "文档" in out and "https://x.com/a" in out

    def test_quote(self):
        assert renderer.parse_markdown("> 引用内容")[0][0] == "quote"

    def test_horizontal_rule(self):
        assert renderer.parse_markdown("---")[0][0] == "hr"

    def test_blank_lines_skipped(self):
        assert renderer.parse_markdown("a\n\n\nb") == [("p", "a"), ("p", "b")]

    def test_empty_input(self):
        assert renderer.parse_markdown("") == []

    def test_unclosed_fence_tolerated(self):
        blocks = renderer.parse_markdown("```py\nx=1")
        assert blocks[0][0] == "code"
        assert "x=1" in blocks[0][1]


# ─────────────────────────────── docx 渲染 ───────────────────────────────

@needs_docx
class TestRenderDocx:
    def _render(self, tmp_path, md=SAMPLE_MD, name="out.docx"):
        res = renderer.render_docx(md, tmp_path / name, title="季度技术报告")
        assert res["status"] == "ok", res
        return res, tmp_path / name

    def test_creates_file(self, tmp_path):
        res, path = self._render(tmp_path)
        assert path.exists()
        assert res["size"] > 0

    def test_uses_real_word_heading_styles(self, tmp_path):
        """落真正的 Heading 样式——Word 才能自动生成目录、支持导航。"""
        from docx import Document
        _, path = self._render(tmp_path)
        doc = Document(path)
        styles = {p.style.name for p in doc.paragraphs if p.text.strip()}
        assert "Title" in styles
        assert "Heading 1" in styles or "Heading 2" in styles

    def test_creates_parent_dir(self, tmp_path):
        _, path = self._render(tmp_path, name="nested/deep/out.docx")
        assert path.exists()

    def test_chinese_font_pinned_in_styles(self, tmp_path):
        """w:eastAsia 必须在样式层钉死，否则跨机打开回退宋体。"""
        import re
        import zipfile
        _, path = self._render(tmp_path)
        with zipfile.ZipFile(path) as z:
            styles_xml = z.read("word/styles.xml").decode("utf-8")
        assert "eastAsia" in styles_xml
        assert re.search(r'w:eastAsia="[^"]*[一-鿿][^"]*"', styles_xml)

    def test_list_bullets(self, tmp_path):
        from docx import Document
        _, path = self._render(tmp_path)
        doc = Document(path)
        assert any(p.style.name == "List Bullet" for p in doc.paragraphs)

    def test_empty_content_errors(self, tmp_path):
        assert renderer.render_docx("", tmp_path / "x.docx")["status"] == "error"

    def test_missing_backend_degrades(self, tmp_path, monkeypatch):
        monkeypatch.setattr(renderer, "HAS_DOCX", False)
        res = renderer.render_docx("x", tmp_path / "x.docx")
        assert res["status"] == "error"
        assert "python-docx" in res["message"]


# ──────────────────────────────── pdf 渲染 ────────────────────────────────

@needs_pdf
class TestRenderPdf:
    def _render(self, tmp_path, md=SAMPLE_MD, name="out.pdf"):
        res = renderer.render_pdf(md, tmp_path / name, title="季度技术报告")
        assert res["status"] == "ok", res
        return res, tmp_path / name

    def test_creates_file(self, tmp_path):
        res, path = self._render(tmp_path)
        assert path.exists()
        assert res["size"] > 0

    @needs_pdftext
    def test_chinese_fully_preserved(self, tmp_path):
        import re

        import pymupdf
        _, path = self._render(tmp_path)
        pdf = pymupdf.open(path)
        text = "".join(p.get_text() for p in pdf)
        cjk = re.compile(r"[\u4e00-\u9fff]")
        assert len(cjk.findall(text)) == len(cjk.findall(SAMPLE_MD))

    @needs_pdftext
    def test_no_mojibake(self, tmp_path):
        import pymupdf
        _, path = self._render(tmp_path)
        text = "".join(p.get_text() for p in pymupdf.open(path))
        assert "�" not in text

    @needs_pdftext
    def test_code_indent_preserved(self, tmp_path):
        """wordWrap="CJK" 会吞行首空格，缩进必须转 &nbsp; 才保得住。"""
        import pymupdf
        _, path = self._render(tmp_path)
        text = "".join(p.get_text() for p in pymupdf.open(path))
        assert "    result = await pipeline.run(doc)" in text

    def test_preserve_indent_helper(self):
        assert renderer._preserve_indent("    x") == "&nbsp;" * 4 + "x"
        assert renderer._preserve_indent("") == ""
        assert renderer._preserve_indent("a<b") == "a&lt;b"

    def test_empty_content_errors(self, tmp_path):
        assert renderer.render_pdf("", tmp_path / "x.pdf")["status"] == "error"

    def test_missing_backend_degrades(self, tmp_path, monkeypatch):
        monkeypatch.setattr(renderer, "HAS_PDF", False)
        res = renderer.render_pdf("x", tmp_path / "x.pdf")
        assert res["status"] == "error"
        assert "reportlab" in res["message"]


# ─────────────────────────────── 统一入口 ───────────────────────────────

class TestRenderDispatch:
    def test_rejects_unknown_format(self, tmp_path):
        res = renderer.render("x", tmp_path / "x.rtf", fmt="rtf")
        assert res["status"] == "error"
        assert "rtf" in res["message"]

    def test_accepts_dot_prefix(self, tmp_path, monkeypatch):
        monkeypatch.setattr(renderer, "HAS_DOCX", False)
        # ".docx" 与 "docx" 等价，不应因前导点被拒
        res = renderer.render("x", tmp_path / "x.docx", fmt=".docx")
        assert "不支持的格式" not in res.get("message", "")

    def test_supported_formats_reflects_env(self):
        avail = renderer.supported_formats()
        assert set(avail) <= {"docx", "pdf", "xlsx", "pptx"}

    def test_clean_inline_variants(self):
        assert renderer.clean_inline("`a` **b** [c](d)") == "a b c (d)"


# ─────────────────────────────── Agent 契约 ───────────────────────────────

def _make_agent(tmp_path, config=None):
    from agents.renderer_agent import RendererAgent
    from pipeline_core.registry import AgentMeta

    class _Bus:
        def subscribe(self, topic, handler=None):
            pass

        def publish(self, topic, sender, payload):
            pass

    class _Reg:
        pass

    meta = AgentMeta(name="renderer", version="1.0")
    cfg = {"formats": ["docx", "pdf"], "output_dir": str(tmp_path)}
    cfg.update(config or {})
    return RendererAgent("renderer", meta, cfg, _Bus(), _Reg())


needs_both = pytest.mark.skipif(
    not (renderer.HAS_DOCX and renderer.HAS_PDF), reason="需要 docx + pdf 后端")


@needs_both
class TestRendererAgent:
    def _msg(self, **payload):
        from pipeline_core.base_agent import Message
        payload.setdefault("task_id", "t1")
        return Message(topic="renderer.render", from_agent="test", payload=payload)

    def test_renders_both_formats(self, tmp_path):
        agent = _make_agent(tmp_path)
        res = agent.handle(self._msg(content=SAMPLE_MD))
        assert res["status"] == "ok"
        assert res["formats"] == ["docx", "pdf"]
        assert Path(res["outputs"]["docx"]).exists()
        assert Path(res["outputs"]["pdf"]).exists()

    def test_basename_from_task_id(self, tmp_path):
        agent = _make_agent(tmp_path)
        res = agent.handle(self._msg(content=SAMPLE_MD, task_id="job_9"))
        assert Path(res["outputs"]["docx"]).name == "job_9.docx"

    def test_basename_from_target_stem(self, tmp_path):
        agent = _make_agent(tmp_path)
        res = agent.handle(self._msg(content=SAMPLE_MD,
                                      target="reports/我的文档.md"))
        assert Path(res["outputs"]["docx"]).name == "我的文档.docx"

    def test_task_id_sanitized(self, tmp_path):
        agent = _make_agent(tmp_path)
        res = agent.handle(self._msg(content=SAMPLE_MD,
                                      task_id="../../evil/path"))
        assert "/" not in Path(res["outputs"]["docx"]).name
        assert ".." not in res["outputs"]["docx"]

    def test_empty_content_errors(self, tmp_path):
        agent = _make_agent(tmp_path)
        assert agent.handle(self._msg(content=""))["status"] == "error"

    def test_unsupported_formats_filtered_by_env(self, tmp_path, monkeypatch):
        monkeypatch.setattr(renderer, "HAS_PDF", False)
        agent = _make_agent(tmp_path, {"formats": ["docx", "pdf"]})
        res = agent.handle(self._msg(content=SAMPLE_MD))
        assert "pdf" not in res["formats"]
        assert res["status"] == "ok"

    def test_all_backends_missing_does_not_crash(self, tmp_path, monkeypatch):
        """后端全缺时如实报错，但应是结构化结果而非抛异常。"""
        monkeypatch.setattr(renderer, "HAS_DOCX", False)
        monkeypatch.setattr(renderer, "HAS_PDF", False)
        agent = _make_agent(tmp_path)
        res = agent.handle(self._msg(content=SAMPLE_MD))
        assert res["status"] == "error"
        assert res["formats"] == []

    def test_parse_formats_string(self, tmp_path):
        agent = _make_agent(tmp_path, {"formats": "docx, pdf"})
        assert agent._formats == ["docx", "pdf"]

    def test_parse_formats_empty_uses_env(self, tmp_path):
        agent = _make_agent(tmp_path, {"formats": []})
        assert agent._formats == renderer.supported_formats()

    def test_drops_unknown_format(self, tmp_path):
        agent = _make_agent(tmp_path, {"formats": ["docx", "rtf"]})
        assert agent._formats == ["docx"]

    def test_resolve_base_name_fallback(self):
        from agents.renderer_agent import RendererAgent
        assert RendererAgent._resolve_base_name({}, "") == "document"


# ────────────────────────── pipeline 声明一致性 ──────────────────────────

class TestPipelineDeclaration:
    def test_docgen_render_parses(self):
        from pipeline_core.scheduler import Scheduler
        sched = Scheduler()
        plan = sched.parse_file(
            str(ROOT / "pipelines" / "docgen-render.yaml"), verify_lock=False)
        assert plan.pipeline_name == "docgen-render"
        names = {n.agent_name for lv in plan.levels for n in lv}
        assert any(n.startswith("renderer") for n in names)

    def test_renderer_is_trusted_agent(self):
        from pathlib import Path

        from pipeline_core.agent_loader import declares_sandbox_trust
        p = Path(__file__).parent.parent / "agents" / "renderer_agent.py"
        assert declares_sandbox_trust(p)

    def test_lockfile_in_sync(self):
        """lockfile 必须与 YAML 一致，否则运行时会被拒绝执行。"""
        from pipeline_core.scheduler import Scheduler
        sched = Scheduler()
        sched.parse_file(
            str(ROOT / "pipelines" / "docgen-render.yaml"), verify_lock=True)


class TestControlCharsAreStripped:
    """抓取来的正文混一个 `\x08` 就会让整条渲染判死（#21 实测）。

    python-docx 在 lxml 层抛 `ValueError: All strings must be XML compatible`，
    reportlab 那条同理；这类字符不可见、不承载内容，剔除才能出厂——但剔了几个
    必须如实回报，正文被动过不该是静默的。
    """

    DIRTY = "# 标题\n\n运行实例 » \x08相关文章\n\n- 项目\x08\n\n代码\t缩进\n换行\r回车\n"

    def test_xml_safe_keeps_tab_cr_lf_and_counts_the_rest(self):
        clean, n = renderer.xml_safe(self.DIRTY)
        assert n == 2 and "\x08" not in clean
        assert "\t" in clean and "\r" in clean and "\n" in clean

    def test_docx_renders_and_reports_the_count(self, tmp_path):
        pytest.importorskip("docx")
        res = renderer.render_docx(self.DIRTY, tmp_path / "d.docx")
        assert res["status"] == "ok", res
        assert res["control_chars_stripped"] == 2

        from docx import Document
        text = "\n".join(p.text for p in Document(res["path"]).paragraphs)
        assert "相关文章" in text and "运行实例" in text, text
        assert "\x08" not in text

    def test_pdf_renders_and_reports_the_count(self, tmp_path):
        pytest.importorskip("reportlab")
        res = renderer.render_pdf(self.DIRTY, tmp_path / "d.pdf")
        assert res["status"] == "ok", res
        assert res["control_chars_stripped"] == 2

    def test_dirty_title_is_cleaned_and_counted(self, tmp_path):
        """标题是另一条入口：docx 写 core_properties 被 suppress 包着，脏标题不报错
        而是静默丢掉文档属性——静默丢属性同样不可接受，所以一并清洗。"""
        pytest.importorskip("docx")
        dirty_title = "季度" + chr(0x0b) + "报告"
        res = renderer.render_docx("# 干净标题\n\n正文内容\n",
                                   tmp_path / "t.docx", title=dirty_title)
        assert res["status"] == "ok" and res["control_chars_stripped"] == 1
        from docx import Document
        assert Document(res["path"]).core_properties.title == "季度报告"

    def test_clean_input_reports_zero(self, tmp_path):
        """没剔就不该虚报：判据要能区分"动过"与"没动过"。"""
        pytest.importorskip("docx")
        res = renderer.render_docx(SAMPLE_MD, tmp_path / "c.docx")
        assert res["status"] == "ok" and res["control_chars_stripped"] == 0


# ────────────────────────────── 表格与图片 ──────────────────────────────

TABLE_MD = """# 数据表

| 指标 | 本季 | 上季 |
| --- | --- | --- |
| 营收 | 120 | 98 |
| 转义 | a\\|b | ok |

正文收尾。
"""


class TestParseTables:
    """表格块识别：认成 table 才走真表格，认不出必须回落段落（不误吞）。"""

    def test_table_block_recognized(self):
        blocks = renderer.parse_markdown("| a | b |\n| --- | --- |\n| 1 | 2 |")
        assert [k for k, _ in blocks] == ["table"]
        assert blocks[0][1] == "| a | b |\n| 1 | 2 |"

    def test_separator_row_dropped(self):
        """分隔行是语法壳，不进数据。"""
        blocks = renderer.parse_markdown("| a |\n| --- |\n| 1 |")
        assert "---" not in blocks[0][1]

    def test_without_separator_falls_back_to_paragraph(self):
        blocks = renderer.parse_markdown("| 单独一行 | 不是表格 |")
        assert [k for k, _ in blocks] == ["p"]

    def test_separator_without_leading_pipe(self):
        """分隔行可以不带首竖线（`--- | ---`）。

        旧实现"先全收再 pop(1)"在这种形态上直接 IndexError——
        收行时按 | 行收，分隔行收不进来，pop 越界。
        """
        blocks = renderer.parse_markdown("| a | b |\n--- | ---\n| 1 | 2 |")
        assert [k for k, _ in blocks] == ["table"]
        assert blocks[0][1] == "| a | b |\n| 1 | 2 |"

    def test_bare_rule_is_not_table_separator(self):
        """单独一行 `---` 是水平线，不是分隔行——表头 + 水平线不得误判成表。"""
        blocks = renderer.parse_markdown("| 看起来像表头 |\n---")
        assert [k for k, _ in blocks] == ["p", "hr"]

    def test_pipe_in_prose_not_swallowed(self):
        """散文里的普通竖线行不是表格——没有分隔行就不认。"""
        blocks = renderer.parse_markdown("a | b | c\nand | more")
        assert [k for k, _ in blocks] == ["p", "p"]

    def test_rows_split_unescaped_pipe_only(self):
        """`\\|` 是单元格内的字面竖线，不能当分隔符切。"""
        assert renderer._table_rows("| a\\|b | c |") == [["a|b", "c"]]

    def test_ragged_rows_padded(self):
        assert renderer._table_rows("| a | b |\n| c |") == [["a", "b"], ["c", ""]]

    def test_empty_middle_cell_kept(self):
        assert renderer._table_rows("| a |  | b |") == [["a", "", "b"]]


class TestTableRendering:
    """端到端：docx 落真 Word 表格（可编辑、可被 Word 识别），pdf 文本保真。"""

    def test_docx_becomes_real_word_table(self, tmp_path):
        pytest.importorskip("docx")
        res = renderer.render_docx(TABLE_MD, tmp_path / "t.docx")
        assert res["status"] == "ok"
        assert res["blocks"]["table"] == 1
        from docx import Document
        doc = Document(res["path"])
        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert len(table.rows) == 3          # 表头 + 2 行数据
        assert [c.text for c in table.rows[0].cells] == ["指标", "本季", "上季"]

    def test_docx_escaped_pipe_restored(self, tmp_path):
        pytest.importorskip("docx")
        res = renderer.render_docx(TABLE_MD, tmp_path / "t.docx")
        from docx import Document
        table = Document(res["path"]).tables[0]
        assert table.rows[2].cells[1].text == "a|b"

    def test_docx_header_bold(self, tmp_path):
        pytest.importorskip("docx")
        res = renderer.render_docx(TABLE_MD, tmp_path / "t.docx")
        from docx import Document
        table = Document(res["path"]).tables[0]
        header_run = table.rows[0].cells[0].paragraphs[0].runs[0]
        body_run = table.rows[1].cells[0].paragraphs[0].runs[0]
        assert header_run.font.bold and not body_run.font.bold

    @needs_pdftext
    def test_pdf_table_text_preserved(self, tmp_path):
        pytest.importorskip("reportlab")
        res = renderer.render_pdf(TABLE_MD, tmp_path / "t.pdf")
        assert res["status"] == "ok"
        import pymupdf
        text = "".join(p.get_text() for p in pymupdf.open(res["path"]))
        assert "指标" in text and "营收" in text
        assert "a|b" in text


class TestCleanInlineImage:
    """图片标记要在链接规则之前吃掉，否则留下 `!` 残渣。"""

    def test_image_marker_stripped(self):
        assert renderer.clean_inline("![截图](https://x/y.png)") == "截图 (https://x/y.png)"

    def test_image_inside_sentence(self):
        out = renderer.clean_inline("见图 ![架构](https://x/a.png) 说明")
        assert out == "见图 架构 (https://x/a.png) 说明"
        assert "!" not in out


# ────────────────────────────── xlsx / pptx ──────────────────────────────

class TestRenderXlsx:
    """结构化子集：每张表一个工作表；无表落「正文」单列（mode 如实回报）。"""

    def test_one_sheet_per_table_named_by_heading(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        res = renderer.render_xlsx(TABLE_MD, tmp_path / "t.xlsx")
        assert res["status"] == "ok"
        assert res["mode"] == "tables" and res["tables"] == 1
        wb = openpyxl.load_workbook(res["path"])
        assert wb.sheetnames == ["数据表"]          # 名取最近的上游标题
        ws = wb["数据表"]
        assert [c.value for c in ws[1]] == ["指标", "本季", "上季"]
        assert ws["A2"].value == "营收"

    def test_header_bold_and_freeze(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        res = renderer.render_xlsx(TABLE_MD, tmp_path / "t.xlsx")
        ws = openpyxl.load_workbook(res["path"]).active
        assert ws["A1"].font.bold is True and ws["B1"].font.bold is True
        assert ws.freeze_panes == "A2"

    def test_escaped_pipe_restored(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        res = renderer.render_xlsx(TABLE_MD, tmp_path / "t.xlsx")
        ws = openpyxl.load_workbook(res["path"]).active
        assert ws["B3"].value == "a|b"

    def test_duplicate_sheet_names_deduped(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        md = ("## 表一\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n"
              "## 表一\n\n| c | d |\n| --- | --- |\n| 3 | 4 |\n")
        res = renderer.render_xlsx(md, tmp_path / "t.xlsx")
        wb = openpyxl.load_workbook(res["path"])
        assert len(wb.sheetnames) == 2
        assert wb.sheetnames[0] != wb.sheetnames[1]

    def test_sheet_name_sanitized(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        md = "## 2026/Q4: 营收[预估]\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n"
        res = renderer.render_xlsx(md, tmp_path / "t.xlsx")
        name = openpyxl.load_workbook(res["path"]).sheetnames[0]
        assert not set(name) & set(":\\/?*[]")
        assert len(name) <= 31

    def test_text_mode_when_no_tables(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        res = renderer.render_xlsx("# 标题\n\n正文一段。\n", tmp_path / "t.xlsx")
        assert res["status"] == "ok" and res["mode"] == "text"
        ws = openpyxl.load_workbook(res["path"]).active
        col = [c.value for c in ws["A"] if c.value is not None]
        assert "标题" in col and "正文一段。" in col

    def test_missing_backend_degrades(self, tmp_path, monkeypatch):
        monkeypatch.setattr(renderer, "HAS_XLSX", False)
        res = renderer.render_xlsx("x", tmp_path / "x.xlsx")
        assert res["status"] == "error" and "openpyxl" in res["message"]

    def test_empty_content_errors(self, tmp_path):
        assert renderer.render_xlsx("", tmp_path / "x.xlsx")["status"] == "error"


class TestRenderPptx:
    """结构化子集：标题起页、正文成要点、表格落真 PowerPoint 表格。"""

    def test_slide_follows_heading(self, tmp_path):
        pytest.importorskip("pptx")
        import pptx
        res = renderer.render_pptx(TABLE_MD, tmp_path / "t.pptx")
        assert res["status"] == "ok" and res["slides"] == 1
        prs = pptx.Presentation(res["path"])
        assert prs.slides[0].shapes.title.text == "数据表"

    def test_bullets_and_real_table(self, tmp_path):
        pytest.importorskip("pptx")
        import pptx
        res = renderer.render_pptx(TABLE_MD, tmp_path / "t.pptx")
        prs = pptx.Presentation(res["path"])
        texts: list[str] = []
        table_cells: list[str] = []
        for shape in prs.slides[0].shapes:
            if shape.has_table:
                tbl = shape.table
                table_cells += [tbl.cell(0, 0).text, tbl.cell(2, 1).text]
            elif shape.has_text_frame:
                texts += [p.text for p in shape.text_frame.paragraphs]
        assert "正文收尾。" in texts                    # 段落成要点
        assert "指标" in table_cells and "a|b" in table_cells   # 真表格 + 转义还原

    def test_each_heading_new_slide(self, tmp_path):
        pytest.importorskip("pptx")
        import pptx
        md = "# 一\n\n- 甲\n\n## 二\n\n段落乙\n\n### 三\n\n结尾\n"
        res = renderer.render_pptx(md, tmp_path / "t.pptx")
        prs = pptx.Presentation(res["path"])
        assert res["slides"] == 3
        assert [s.shapes.title.text for s in prs.slides] == ["一", "二", "三"]

    def test_oversized_table_truncated_with_note(self, tmp_path):
        """超限就地截断并把原表规模写在页内——溢出在 pptx 里是看不见的。"""
        pytest.importorskip("pptx")
        import pptx
        rows = "\n".join(f"| r{i} | v{i} |" for i in range(20))
        md = f"# 大表\n\n| 甲 | 乙 |\n| --- | --- |\n{rows}\n"
        res = renderer.render_pptx(md, tmp_path / "t.pptx")
        tbl = next(s for s in pptx.Presentation(res["path"]).slides[0].shapes
                   if s.has_table).table
        assert len(tbl.rows) == renderer._PPTX_MAX_TABLE_ROWS + 2   # 表头+12行+注记行
        note = tbl.cell(len(tbl.rows) - 1, 0).text
        assert "原表 21 行" in note

    def test_chinese_font_pinned(self, tmp_path):
        pytest.importorskip("pptx")
        import zipfile
        res = renderer.render_pptx(TABLE_MD, tmp_path / "t.pptx")
        with zipfile.ZipFile(res["path"]) as z:
            xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
        assert 'typeface="微软雅黑"' in xml and "a:ea" in xml

    def test_dispatch_by_format_name(self, tmp_path):
        pytest.importorskip("pptx")
        res = renderer.render("# 标题\n\n内容\n", tmp_path / "d.pptx", fmt="pptx")
        assert res["status"] == "ok" and res["format"] == "pptx"

    def test_missing_backend_degrades(self, tmp_path, monkeypatch):
        monkeypatch.setattr(renderer, "HAS_PPTX", False)
        res = renderer.render_pptx("x", tmp_path / "x.pptx")
        assert res["status"] == "error" and "python-pptx" in res["message"]

    def test_empty_content_errors(self, tmp_path):
        assert renderer.render_pptx("", tmp_path / "x.pptx")["status"] == "error"
