"""校验真实流水线产物（output/render_e2e.*）的保真度。"""

import re
from pathlib import Path

import pymupdf
from docx import Document

OUT = Path(__file__).resolve().parent.parent / "output"

CJK = re.compile(r"[\u4e00-\u9fff]")


def main() -> None:
    md = (OUT / "render_e2e.md").read_text(encoding="utf-8")

    pdf = pymupdf.open(OUT / "render_e2e.pdf")
    ptxt = "".join(p.get_text() for p in pdf)
    print("PDF 页数:", pdf.page_count)
    print("中文字符:", len(CJK.findall(md)), "->", len(CJK.findall(ptxt)),
          f"({len(CJK.findall(ptxt)) / max(len(CJK.findall(md)), 1):.0%} 保留)")
    print("PDF 无乱码豆腐块:", "�" not in ptxt)

    doc = Document(OUT / "render_e2e.docx")
    heads = [p.text for p in doc.paragraphs
             if p.style.name.startswith("Heading") or p.style.name == "Title"]
    paras = [p for p in doc.paragraphs if p.text.strip()]
    print("docx 非空段落:", len(paras))
    print("docx Word 标题段落:", len(heads))
    for h in heads[:6]:
        print("   -", h)


if __name__ == "__main__":
    main()
